"""20261030_졸업논문.md -> docx (한국기술교육대학교 학위논문 양식).

thesis_format.py 의 판형·스타일·구역·캡션 후처리를 그대로 재사용하되,
수식 번호는 md 의 \\tag{장-번호} 라벨을 문서 순서대로 그대로 사용한다.

실행: python GRADUATION/tools/thesis_docx_20261030.py
"""
import re
import sys
from pathlib import Path

import pypandoc
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
import thesis_format as tf  # noqa: E402

ROOT = tf.ROOT
SRC_MD = ROOT / "20261030_졸업논문.md"
OUT_DOCX = ROOT / "20261030_졸업논문.docx"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def strip_tags(md):
    """\\tag{label} 을 제거하고 라벨 목록을 문서 순서대로 돌려준다."""
    labels = re.findall(r"\\tag\{([^}]+)\}", md)
    md = re.sub(r"\s*\\tag\{[^}]+\}\$\$", "$$", md)
    return md, labels


def postprocess(docx_path, labels):
    doc = Document(str(docx_path))
    for sec in doc.sections:
        sec.page_width, sec.page_height = tf.PAGE_W, tf.PAGE_H
        sec.left_margin, sec.right_margin = tf.MARGIN_L, tf.MARGIN_R
        sec.top_margin, sec.bottom_margin = tf.MARGIN_T, tf.MARGIN_B
        sec.footer_distance = Mm(12)
        sec.start_type = WD_SECTION.NEW_PAGE
    secs = doc.sections
    if len(secs) != 3:
        raise RuntimeError(f"expected 3 sections, got {len(secs)}")
    tf.page_footer(secs[0], None)
    tf.page_footer(secs[1], "lowerRoman")
    tf.page_footer(secs[2], "decimal")
    paras = doc.paragraphs
    for k, p in enumerate(paras[:-1]):
        if p._p.find(qn("w:pPr")) is not None and p._p.pPr.find(qn("w:sectPr")) is not None:
            paras[k + 1].paragraph_format.page_break_before = False

    eq_idx = 0
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt in ("@@TOC@@", "@@TOT@@", "@@TOF@@"):
            for r in list(p.runs):
                r._r.getparent().remove(r._r)
            instr = {"@@TOC@@": 'TOC \\o "1-3" \\h \\z \\u',
                     "@@TOT@@": 'TOC \\h \\z \\t "표캡션,1"',
                     "@@TOF@@": 'TOC \\h \\z \\t "그림캡션,1"'}[txt]
            p.style = doc.styles["Normal"]
            tf.add_field(p, instr, "목록을 갱신하려면 F9를 누르십시오.")
            continue
        if re.match(r"^<표\d+-\d+>", txt):
            p.style = doc.styles["표캡션"]
            continue
        if re.match(r"^\[그림 \d+-\d+\]", txt):
            p.style = doc.styles["그림캡션"]
            continue
        omp = p._p.find(f"{{{M_NS}}}oMathPara")
        if omp is not None:
            label = labels[eq_idx] if eq_idx < len(labels) else "?"
            eq_idx += 1
            p.style = doc.styles["수식"]
            maths = omp.findall(f"{{{M_NS}}}oMath")
            idx = list(p._p).index(omp)
            p._p.remove(omp)

            def tab_run():
                r = OxmlElement("w:r")
                r.append(OxmlElement("w:tab"))
                return r
            seq = [tab_run()] + maths + [tab_run()]
            num = OxmlElement("w:r")
            t = OxmlElement("w:t")
            t.text = f"({label})"
            num.append(t)
            seq.append(num)
            for k, el in enumerate(seq):
                p._p.insert(idx + k, el)
            tabs = p.paragraph_format.tab_stops
            tabs.add_tab_stop(tf.TEXT_W // 2, WD_TAB_ALIGNMENT.CENTER)
            tabs.add_tab_stop(tf.TEXT_W, WD_TAB_ALIGNMENT.RIGHT)

    for k, tbl in enumerate(doc.tables):
        if k == 0:  # 책등 표
            continue
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tf.table_borders(tbl)
        rows = tbl.rows
        for ri, row in enumerate(rows):
            # 행 내부 분할 금지 + 마지막 행 전까지 다음 행과 묶어 표가 쪽 경계에서 쪼개지지 않게
            trPr = row._tr.get_or_add_trPr()
            if trPr.find(qn("w:cantSplit")) is None:
                trPr.append(OxmlElement("w:cantSplit"))
            for cell in row.cells:
                for cp in cell.paragraphs:
                    cp.paragraph_format.first_line_indent = 0
                    cp.paragraph_format.line_spacing = 1.2
                    cp.paragraph_format.keep_with_next = ri < len(rows) - 1
                    for r in cp.runs:
                        r.font.size = Pt(9.5)

    # 표 직후 문단은 표 테두리에 밀착하므로 위 간격 6pt
    for k, tbl in enumerate(doc.tables):
        if k == 0:
            continue
        nxt = tbl._tbl.getnext()
        if nxt is not None and nxt.tag == qn("w:p"):
            pPr = nxt.get_or_add_pPr()
            sp = pPr.find(qn("w:spacing"))
            if sp is None:
                sp = OxmlElement("w:spacing")
                pPr.append(sp)
            sp.set(qn("w:before"), "120")

    settings = doc.settings.element
    uf = OxmlElement("w:updateFields")
    uf.set(qn("w:val"), "true")
    settings.append(uf)
    doc.save(str(docx_path))
    return eq_idx


def main():
    md = SRC_MD.read_text(encoding="utf-8")
    md, labels = strip_tags(md)
    print("md:", SRC_MD.name, "equations:", len(labels))

    tmp = ROOT / "tools" / "_build"
    tmp.mkdir(exist_ok=True)
    ref = tmp / "reference.docx"
    tf.build_reference(ref)
    pin = tmp / "pandoc_input_20261030.md"
    pin.write_text(tf.pandoc_input(md), encoding="utf-8")
    pypandoc.convert_file(
        str(pin), "docx", format="markdown-implicit_figures+raw_attribute",
        outputfile=str(OUT_DOCX),
        extra_args=[f"--reference-doc={ref}", f"--resource-path={ROOT}"])
    n_eq = postprocess(OUT_DOCX, labels)
    print("docx:", OUT_DOCX.name, "equations numbered:", n_eq, "labels:", labels[0], "...", labels[-1])
    if n_eq != len(labels):
        sys.exit(f"equation count mismatch: md={len(labels)} docx={n_eq}")


if __name__ == "__main__":
    main()
